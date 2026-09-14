"""PDF and PPTX ingestion: text, titles, tables, notes, charts, images, snapshot facts."""

from __future__ import annotations

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from app.config import LOGO_PATH
from app.ingestion.extractor import extract_company_snapshot, load_deck
from app.models.schemas import Classification


# --- PDF ------------------------------------------------------------------------

def test_pdf_pages_titles_and_text(risky_pdf, rules_settings):
    deck = load_deck(risky_pdf, rules_settings)
    assert deck.file_type == "pdf"
    assert deck.slide_count == 9
    assert len(deck.sha256) == 64
    assert deck.slide(1).title == "QuickMind AI"
    assert deck.slide(4).title == "Technology"
    assert "offshore development agency" in deck.slide(4).full_text
    assert deck.slide(10) is None


def test_pdf_table_extraction(strong_pdf, rules_settings):
    deck = load_deck(strong_pdf, rules_settings)
    tables = deck.slide(4).tables
    assert tables, "ruled table on the architecture slide should be extracted"
    cells = [cell for table in tables for row in table for cell in row]
    assert any("Kafka event streaming" in cell for cell in cells)


def test_pdf_snapshot_facts_keep_slide_numbers(risky_pdf, rules_settings):
    snapshot = extract_company_snapshot(load_deck(risky_pdf, rules_settings))
    assert snapshot.name.value == "QuickMind AI" and snapshot.name.slides == [1]
    assert snapshot.revenue.value == "$180K revenue"
    assert snapshot.revenue.slides == [7]
    assert snapshot.raise_amount.value.startswith("$3M")
    assert snapshot.customer_concentration.value == "Largest customer 70% of revenue"
    assert snapshot.gross_margin.classification == Classification.NOT_PROVIDED


# --- PPTX -----------------------------------------------------------------------

def test_pptx_titles_tables_groups_and_notes(strong_pptx, rules_settings):
    deck = load_deck(strong_pptx, rules_settings)
    assert deck.file_type == "pptx"
    assert deck.slide_count == 9
    assert deck.slide(5).title == "Engineering & Delivery"
    assert any("Kafka event streaming" in cell for row in deck.slide(4).tables[0] for cell in row)
    assert "Engineering leadership: CTO, VP Engineering" in deck.slide(5).full_text  # grouped shape
    assert "Headcount figures" in deck.slide(5).notes
    assert deck.metadata_title == "Helix Grid Systems investor presentation"


def test_pptx_snapshot(strong_pptx, rules_settings):
    snapshot = extract_company_snapshot(load_deck(strong_pptx, rules_settings))
    assert snapshot.revenue.value == "$14.2M ARR" and snapshot.revenue.slides == [3]
    assert snapshot.growth.value == "118% YoY growth"
    assert snapshot.gross_margin.value == "78%"
    assert snapshot.raise_amount.value == "$40M Series B"
    assert snapshot.funding_to_date.value == "$31M raised to date"
    assert snapshot.engineering_team.value == "46 engineers"
    assert snapshot.sector.classification == Classification.INFERENCE


def test_pptx_chart_and_picture_extraction(tmp_path, rules_settings):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Revenue Growth"
    data = CategoryChartData()
    data.categories = ["2023", "2024", "2025"]
    data.add_series("ARR ($M)", (2.1, 6.5, 14.2))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4),
                                   data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "ARR by year"
    slide.shapes.add_picture(str(LOGO_PATH), Inches(7.5), Inches(1.5))
    path = tmp_path / "chart.pptx"
    presentation.save(str(path))

    parsed = load_deck(path, rules_settings).slide(1)
    assert "ARR by year" in parsed.full_text
    assert "2024" in parsed.full_text and "ARR ($M)" in parsed.full_text
    assert parsed.image_count == 1


def test_projections_are_not_reported_as_revenue(make_deck, rules_settings):
    path = make_deck(("Financial Plan", ["Projected revenue of $50M by 2029."]))
    snapshot = extract_company_snapshot(load_deck(path, rules_settings))
    assert snapshot.revenue.classification == Classification.NOT_PROVIDED
