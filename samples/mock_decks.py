"""Mock pitch decks used as sample inputs and test fixtures.

Two fictional companies, each renderable as PPTX or PDF:

* Helix Grid Systems — technically strong, clean IP (Series B).
* QuickMind AI — contractor-built code, copyleft open source, founder dependency,
  weak documentation and unresolved university IP (seed extension).

All names, people and figures are invented for testing.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class SlideSpec:
    title: str
    bullets: list[str]
    table: list[list[str]] | None = None
    group: list[str] = field(default_factory=list)
    notes: str = ""


@dataclass
class DeckSpec:
    company: str
    slides: list[SlideSpec]


STRONG_COMPANY = DeckSpec("Helix Grid Systems", [
    SlideSpec("Helix Grid Systems", ["Grid-edge orchestration software for electric utilities",
                                     "Series B Investor Presentation | Confidential"]),
    SlideSpec("The Problem", [
        "Utilities must balance four times more distributed energy resources with dispatch tools built for one-way grids.",
        "Grid operators lose an estimated $9B annually to curtailment and manual balancing."]),
    SlideSpec("Product & Traction", [
        "HelixOS is in production at 38 utilities and generally available since 2023.",
        "$14.2M ARR with 118% YoY growth and 78% gross margin.",
        "Largest customer represents 9% of revenue; net revenue retention of 131%.",
        "Processes 2.1 billion telemetry events per day across 1.4 million grid devices."]),
    SlideSpec("Platform Architecture", [
        "Cloud-native microservices on Kubernetes, deployed multi-region on AWS and Azure with active-active failover.",
        "Public REST and gRPC APIs with SDKs for utility integrations.",
        "Infrastructure as code with Terraform; observability with Datadog and OpenTelemetry.",
        "Autoscaling event pipeline load-tested to 10x current peak volume."],
        table=[["Layer", "Technology"], ["Ingestion", "Kafka event streaming"],
               ["Forecasting", "Proprietary forecasting models built in-house"], ["Data", "PostgreSQL and TimescaleDB"]]),
    SlideSpec("Engineering & Delivery", [
        "CTO Dr. Maya Chen led grid software engineering at Siemens for nine years.",
        "46 full-time in-house engineers across Denver and Toronto.",
        "CI/CD with GitHub Actions; 84% automated test coverage; weekly releases; mandatory code reviews.",
        "99.95% uptime SLA over the trailing 24 months.",
        "Incident response with on-call rotation and blameless post-mortems.",
        "Architecture documentation, runbooks and decision records maintained for every service."],
        group=["Engineering leadership: CTO, VP Engineering, Director of SRE"],
        notes="Headcount figures as of the most recent quarter."),
    SlideSpec("Security & Compliance", [
        "SOC 2 Type II attested annually and ISO 27001 certified.",
        "Annual third-party penetration testing and a public bug bounty program.",
        "Data encrypted with AES-256 at rest and TLS 1.3 in transit; SSO, MFA and role-based access control.",
        "Disaster recovery tested quarterly with 15-minute RPO and 1-hour RTO."]),
    SlideSpec("Intellectual Property", [
        "7 issued U.S. patents and 5 PCT applications covering forecasting and dispatch methods.",
        "All patents are assigned to Helix Grid Systems, Inc.",
        "All employees and contractors have signed invention assignment agreements.",
        "Proprietary dataset of 11 billion labeled grid telemetry records.",
        "Open-source license scanning with FOSSA in CI; no copyleft components in distributed code.",
        "Independent IP and code audit completed in 2025."]),
    SlideSpec("Market & Competition", [
        "Embedded in utility control-room workflows as the system of record for DER dispatch.",
        "High switching costs from integrations with SCADA and ADMS platforms."]),
    SlideSpec("The Round", [
        "Raising $40M Series B to expand into European utilities.",
        "Raised $31M to date from climate and infrastructure investors.",
        "Use of funds: EU expansion, FedRAMP authorization planned for 2027, and grid-forecasting R&D."]),
])

RISKY_COMPANY = DeckSpec("QuickMind AI", [
    SlideSpec("QuickMind AI", ["Generative AI copilots for insurance claims teams",
                               "Seed Extension | Investor Presentation"]),
    SlideSpec("The Problem", [
        "Insurance carriers spend 30% of claims cost on manual document review.",
        "Adjusters juggle PDFs, emails and aging claims systems."]),
    SlideSpec("Product", [
        "Our MVP is in beta with 3 pilot customers at regional insurers.",
        "QuickMind is built on the OpenAI GPT-4 API with a prompt layer on top of the model.",
        "Pilots show 40% faster claim triage."]),
    SlideSpec("Technology", [
        "The platform was built by an offshore development agency in 2024.",
        "Our technical co-founder wrote the core claims-parsing engine and is our sole developer.",
        "Monolithic Python application running on a single AWS server.",
        "Uses GPL-licensed open-source libraries and scraped public claims data to improve prompts.",
        "Manual testing before each release; little formal documentation."]),
    SlideSpec("Team", [
        "CEO Jordan Blake, former claims adjuster.",
        "Hiring a CTO after this round.",
        "Engineering: 1 in-house developer plus freelance developers from Upwork."]),
    SlideSpec("Intellectual Property", [
        "Core algorithms were developed during our founder's PhD research at State University.",
        "We plan to file patents after the seed extension.",
        "IP assignment from the university is still being finalized."]),
    SlideSpec("Traction", [
        "3 pilot customers; largest pilot represents 70% of revenue.",
        "$180K revenue in the last twelve months.",
        "Each pilot required a custom integration for each customer, and we maintain a separate code branch for each customer."]),
    SlideSpec("Security", [
        "SOC 2 Type II audit planned for next year.",
        "No formal security program yet; data stored in one region."]),
    SlideSpec("The Ask", [
        "Raising $3M seed extension to hire a CTO and rebuild the platform.",
        "Pre-money valuation of $18M."]),
])


def build_pptx(spec: DeckSpec, path: str | Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    path = Path(path)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    presentation.core_properties.title = f"{spec.company} investor presentation"
    layout = presentation.slide_layouts[5]  # "Title Only"
    for slide_spec in spec.slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = slide_spec.title
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(12.1), Inches(3.0))
        frame = box.text_frame
        frame.word_wrap = True
        for index, bullet in enumerate(slide_spec.bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(16)
        top = Inches(4.7)
        if slide_spec.table:
            rows, cols = len(slide_spec.table), len(slide_spec.table[0])
            shape = slide.shapes.add_table(rows, cols, Inches(0.6), top, Inches(12.1), Inches(0.35 * rows))
            for r, row in enumerate(slide_spec.table):
                for c, value in enumerate(row):
                    shape.table.cell(r, c).text = value
            top = Inches(4.7 + 0.4 * rows)
        if slide_spec.group:
            group = slide.shapes.add_group_shape()
            for offset, text in enumerate(slide_spec.group):
                child = group.shapes.add_textbox(Inches(0.6 + offset * 4), top, Inches(3.8), Inches(0.5))
                child.text_frame.text = text
        if slide_spec.notes:
            slide.notes_slide.notes_text_frame.text = slide_spec.notes
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))
    return path


def build_pdf(spec: DeckSpec, path: str | Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Table, TableStyle

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    width, height = landscape(letter)
    pdf = canvas.Canvas(str(path), pagesize=(width, height))
    pdf.setTitle(f"{spec.company} investor presentation")
    for slide_spec in spec.slides:
        pdf.setFont("Helvetica-Bold", 26)
        pdf.drawString(50, height - 72, slide_spec.title)
        y = height - 115
        pdf.setFont("Helvetica", 13)
        for bullet in slide_spec.bullets + slide_spec.group:
            for line in simpleSplit(bullet, "Helvetica", 13, width - 110):
                pdf.drawString(60, y, line)
                y -= 17
            y -= 7
        if slide_spec.table:
            table = Table(slide_spec.table, colWidths=[180, width - 290])
            table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.7, colors.black),
                                       ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                                       ("FONTSIZE", (0, 0), (-1, -1), 11)]))
            _, table_height = table.wrapOn(pdf, width - 100, y)
            table.drawOn(pdf, 60, y - table_height - 8)
        pdf.showPage()
    pdf.save()
    return path
