"""PowerPoint parsing: PPTX via python-pptx, legacy PPT via a LibreOffice conversion.

Text is collected from placeholders, text boxes, grouped shapes, tables, chart
titles/categories, SmartArt diagram data and speaker notes, in reading order,
so architecture diagrams built from shapes still contribute evidence.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from app.ingestion.document import Slide
from app.ingestion.errors import DeckIngestionError, DeckTooLargeError, MalformedDeckError, UnsupportedDeckError
from app.utils.logging import get_logger
from app.utils.text import clean

log = get_logger("ingestion.ppt")

_DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def parse_pptx(path: Path, max_slides: int) -> tuple[list[Slide], str, list[str]]:
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
    except Exception as exc:  # PackageNotFoundError, BadZipFile, KeyError, XMLSyntaxError...
        log.error("PPTX %s could not be opened: %s: %s", path.name, type(exc).__name__, exc)
        raise MalformedDeckError(
            f"{path.name} could not be read as a PowerPoint file. It may be corrupt, "
            "password-protected, or not a real .pptx."
        ) from exc

    slide_list = list(presentation.slides)
    if len(slide_list) > max_slides:
        raise DeckTooLargeError(f"{path.name} has {len(slide_list)} slides; the limit is {max_slides}.")

    warnings: list[str] = []
    slides: list[Slide] = []
    for number, pptx_slide in enumerate(slide_list, 1):
        try:
            slides.append(_parse_slide(pptx_slide, number))
        except Exception as exc:
            log.warning("Slide %d of %s could not be parsed: %s", number, path.name, exc)
            slides.append(Slide(number=number, warnings=[f"S{number}: slide could not be parsed ({type(exc).__name__})."]))

    try:
        title = clean(presentation.core_properties.title or "")
    except Exception:
        title = ""
    return slides, title, warnings


def parse_ppt(path: Path, max_slides: int, soffice_path: str | None = None) -> tuple[list[Slide], str, list[str]]:
    """Convert a legacy binary .ppt to .pptx with LibreOffice, then parse it."""
    soffice = soffice_path or shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not Path(soffice).exists() and not shutil.which(soffice):
        raise UnsupportedDeckError(
            f"{path.name} is a legacy .ppt file. Converting it requires LibreOffice "
            "(set TDIP_SOFFICE_PATH), or re-save the deck as .pptx or PDF and upload again."
        )
    with tempfile.TemporaryDirectory(prefix="tdip_ppt_") as tmp:
        try:
            completed = subprocess.run(
                [soffice, "--headless", "--convert-to", "pptx", "--outdir", tmp, str(path)],
                capture_output=True, timeout=180, check=False,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            log.error("LibreOffice conversion failed for %s: %s", path.name, exc)
            raise MalformedDeckError(f"{path.name} could not be converted from .ppt.") from exc
        converted = Path(tmp) / f"{path.stem}.pptx"
        if completed.returncode != 0 or not converted.exists():
            log.error("LibreOffice returned %s for %s: %s", completed.returncode, path.name,
                      completed.stderr.decode(errors="ignore")[:500])
            raise MalformedDeckError(f"{path.name} could not be converted from .ppt; it may be corrupt.")
        slides, title, warnings = parse_pptx(converted, max_slides)
        warnings.append("Legacy .ppt converted to .pptx with LibreOffice before extraction.")
        return slides, title, warnings


def _ordered(shapes) -> list:
    def key(shape):
        return (getattr(shape, "top", 0) or 0, getattr(shape, "left", 0) or 0)

    try:
        return sorted(list(shapes), key=key)
    except Exception:
        return list(shapes)


def _parse_slide(pptx_slide, number: int) -> Slide:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    slide = Slide(number=number)
    title_shape = None
    try:
        title_shape = pptx_slide.shapes.title
    except Exception:
        title_shape = None
    title_id = getattr(title_shape, "shape_id", None)
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        slide.title = clean(title_shape.text_frame.text)

    texts: list[str] = []

    def collect(shape) -> None:
        try:
            shape_type = shape.shape_type
        except Exception:
            shape_type = None
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in _ordered(shape.shapes):
                collect(child)
            return
        if shape_type == MSO_SHAPE_TYPE.PICTURE or type(shape).__name__ in {"Picture", "PlaceholderPicture"}:
            slide.image_count += 1
        if getattr(shape, "has_text_frame", False) and shape.shape_id != title_id:
            value = clean(shape.text_frame.text)
            if value:
                texts.append(value)
        if getattr(shape, "has_table", False):
            rows = [[clean(cell.text) for cell in row.cells] for row in shape.table.rows]
            rows = [row for row in rows if any(row)]
            if rows:
                slide.tables.append(rows)
        if getattr(shape, "has_chart", False):
            slide.chart_text.extend(_chart_text(shape.chart))
        diagram = _smartart_text(shape, pptx_slide)
        if diagram:
            slide.diagram_text.append(diagram)

    for shape in _ordered(pptx_slide.shapes):
        try:
            collect(shape)
        except Exception as exc:
            log.debug("Shape skipped on slide %d: %s", number, exc)
            slide.warnings.append(f"S{number}: a shape could not be read ({type(exc).__name__}).")

    slide.text = clean("\n".join(texts))
    if not slide.title and texts:
        first = texts[0].splitlines()[0]
        if len(first) <= 90:
            slide.title = first

    try:
        if pptx_slide.has_notes_slide and pptx_slide.notes_slide.notes_text_frame is not None:
            slide.notes = clean(pptx_slide.notes_slide.notes_text_frame.text)
    except Exception as exc:
        log.debug("Notes unreadable on slide %d: %s", number, exc)
    return slide


def _chart_text(chart) -> list[str]:
    parts: list[str] = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            parts.append(clean(chart.chart_title.text_frame.text))
    except Exception:
        pass
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories if c is not None]
        if categories:
            parts.append("Chart categories: " + ", ".join(categories[:20]))
        names = [s.name for s in plot.series if s.name]
        if names:
            parts.append("Chart series: " + ", ".join(names[:10]))
    except Exception:
        pass
    return [p for p in parts if p]


def _smartart_text(shape, pptx_slide) -> str:
    """Read the text of a SmartArt diagram from its separate diagram-data part."""
    try:
        element = shape._element
        rel_ids = element.find(f".//{{{_DGM_NS}}}relIds")
        if rel_ids is None:
            return ""
        r_id = rel_ids.get(f"{{{_REL_NS}}}dm")
        part = pptx_slide.part.related_part(r_id)
        from lxml import etree

        root = etree.fromstring(part.blob)
        return clean(" ".join(t.text for t in root.iter(f"{{{_A_NS}}}t") if t.text))
    except Exception:
        return ""


def looks_like_pptx(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as archive:
            return "ppt/presentation.xml" in archive.namelist()
    except (zipfile.BadZipFile, OSError):
        return False


__all__ = ["parse_pptx", "parse_ppt", "looks_like_pptx", "DeckIngestionError"]
